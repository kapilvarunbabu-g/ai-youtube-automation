from datetime import datetime, timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .common import ROOT

XLSX = ROOT / "content_log.xlsx"
SHEETS = {
    "Research": ["Run Date", "Kind", "Query", "Title", "Source/Channel", "URL", "Published At", "Views", "Views/Hour", "Likes", "Comments", "Engagement Rate", "Freshness", "Topic Overlap", "Status"],
    "Daily Picks": ["Run Date", "Topic", "Title", "Hook", "Confidence", "Trend Evidence", "Source Links", "Decision"],
    "Video Queue": ["Run Date", "Title", "Script", "Video Prompt", "Description", "Hashtags", "Status"],
    "Published": ["Run Date", "Title", "Video ID", "YouTube URL", "Status"],
}

def style_sheet(ws):
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    ws.auto_filter.ref = ws.dimensions

def ensure_workbook():
    if XLSX.exists():
        workbook = load_workbook(XLSX)
    else:
        workbook = Workbook()
        workbook.remove(workbook.active)
    for name, headers in SHEETS.items():
        if name not in workbook.sheetnames:
            ws = workbook.create_sheet(name)
            ws.append(headers)
            style_sheet(ws)
    workbook.save(XLSX)

def append_row(ws, values):
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)

def autosize(ws):
    for col in range(1, ws.max_column + 1):
        longest = 0
        for cell in ws[get_column_letter(col)]:
            if cell.value is not None:
                longest = max(longest, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col)].width = min(55, max(12, longest + 2))

def update_workbook(research_rows, selected_story, published):
    ensure_workbook()
    workbook = load_workbook(XLSX)
    run_date = datetime.now(timezone.utc).date().isoformat()

    if research_rows:
        ws = workbook["Research"]
        existing_urls = {ws.cell(row, 6).value for row in range(2, ws.max_row + 1)}
        for item in research_rows:
            if item.get("url") in existing_urls:
                continue
            append_row(ws, [
                run_date, item.get("kind"), item.get("query"), item.get("title"),
                item.get("source") or item.get("channel"), item.get("url"), item.get("published_at"),
                item.get("views"), item.get("views_per_hour"), item.get("likes"), item.get("comments"),
                item.get("engagement_rate"), item.get("freshness_score"), item.get("topic_overlap"), "researched",
            ])
        autosize(ws)

    if selected_story:
        picks = workbook["Daily Picks"]
        append_row(picks, [
            run_date,
            selected_story.get("topic"), selected_story.get("title"), selected_story.get("hook"),
            selected_story.get("confidence"), "See research.json and source links",
            "\n".join(s.get("url", "") for s in selected_story.get("sources", [])), "selected",
        ])
        queue = workbook["Video Queue"]
        append_row(queue, [
            run_date, selected_story.get("title"), selected_story.get("script"),
            selected_story.get("video_prompt"), selected_story.get("description"),
            " ".join(selected_story.get("hashtags", [])), "rendered",
        ])
        autosize(picks)
        autosize(queue)

    if published:
        ws = workbook["Published"]
        append_row(ws, [
            run_date, selected_story.get("title"), published["video_id"], published["url"], "uploaded",
        ])
        autosize(ws)

    workbook.save(XLSX)
